import os
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from pptx import Presentation

REPORT_DIR = "reports"
os.makedirs(REPORT_DIR, exist_ok=True)


class ReportAgent:

    def __init__(self):
        pass

    # -------------------------
    # DOCX Report
    # -------------------------
    def generate_docx(self, title, analysis, citations, suffix=None):
        import uuid
        if not suffix:
            suffix = uuid.uuid4().hex[:8]

        doc = Document()

        doc.add_heading(title, level=1)

        doc.add_heading("Analysis", level=2)
        doc.add_paragraph(str(analysis))

        doc.add_heading("Citations", level=2)
        doc.add_paragraph(str(citations))

        filename = f"report_{suffix}.docx"
        path = os.path.join(REPORT_DIR, filename)
        doc.save(path)

        return path

    # -------------------------
    # PDF Report
    # -------------------------
    def generate_pdf(self, title, analysis, citations, suffix=None):
        import uuid
        if not suffix:
            suffix = uuid.uuid4().hex[:8]

        filename = f"report_{suffix}.pdf"
        path = os.path.join(REPORT_DIR, filename)
        doc = SimpleDocTemplate(path)

        styles = getSampleStyleSheet()
        content = []

        content.append(Paragraph(title, styles["Title"]))
        content.append(Spacer(1, 12))

        content.append(Paragraph(str(analysis), styles["BodyText"]))
        content.append(Spacer(1, 12))

        content.append(Paragraph(str(citations), styles["BodyText"]))

        doc.build(content)

        return path

    # -------------------------
    # PPT Report
    # -------------------------
    def generate_ppt(self, title, analysis, suffix=None):
        import uuid
        if not suffix:
            suffix = uuid.uuid4().hex[:8]

        prs = Presentation()

        # Slide 1: Title
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title

        # Slide 2: Analysis
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Analysis"
        slide.placeholders[1].text = str(analysis)

        filename = f"report_{suffix}.pptx"
        path = os.path.join(REPORT_DIR, filename)
        prs.save(path)

        return path


# -------------------------
# TEST
# -------------------------
if __name__ == "__main__":

    agent = ReportAgent()

    analysis = {
        "model": "CNN",
        "accuracy": "96%",
        "dataset": "Agriculture dataset"
    }

    citations = "Smith et al. (2024). Deep Learning for Agriculture."

    print(agent.generate_docx("Research Report", analysis, citations))
    print(agent.generate_pdf("Research Report", analysis, citations))
    print(agent.generate_ppt("Research Report", analysis))